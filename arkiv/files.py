from __future__ import annotations

import base64
import io
import os
import re
import shutil
from pathlib import Path

from arkiv.paths import APP_DIR

TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".jsonl", ".xml",
    ".html", ".htm", ".css", ".js", ".mjs", ".ts", ".tsx", ".jsx", ".py",
    ".rb", ".go", ".rs", ".java", ".c", ".cpp", ".h", ".hpp", ".cs", ".php",
    ".sh", ".zsh", ".yml", ".yaml", ".toml", ".ini", ".cfg", ".log", ".rtf",
    ".sql", ".swift", ".kt", ".lua", ".r", ".tex", ".rst", ".svg",
}
IMAGE_EXTENSIONS = {
    ".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff", ".heic", ".heif",
}
PDF_EXTENSIONS = {".pdf"}
DOCX_EXTENSIONS = {".docx"}
SHEET_EXTENSIONS = {".xlsx", ".xlsm"}
ALWAYS_SKIP_DIRS = {".git", ".svn", "node_modules", "__pycache__", ".venv", ".idea", ".cursor"}
IGNORED_NAMES = {".ds_store", ".localized", "thumbs.db", "desktop.ini"}
TEMP_SUFFIXES = (".crdownload", ".download", ".part", ".tmp", ".temp", ".partial")
CONTENT_LIMIT = 30 * 1024 * 1024
IMAGE_LIMIT = 15 * 1024 * 1024


PARKING_CATEGORIES = {"needs review", "review", "unsorted", "misc", "miscellaneous", "other", "unknown"}


def is_parking_category(name: str) -> bool:
    return str(name or "").casefold().strip() in PARKING_CATEGORIES


def clean_category(name: str) -> str:
    text = str(name or "").replace("..", " ")
    text = re.sub(r'[\\/:*?"<>|]+', " ", text)
    text = re.sub(r"[\x00-\x1f]", " ", text)
    text = re.sub(r"\s+", " ", text).strip(" .")
    if not text or text.casefold() in {"none", "null", "n/a", "undefined", "category"}:
        return "Documents"
    if is_parking_category(text):
        return "Documents"
    words = text.split(" ")
    if len(words) > 4:
        text = " ".join(words[:4])
    if len(text) > 40:
        text = text[:40].rstrip()
    return text or "Documents"


def clean_filename(name: str, extension: str) -> str:
    """Build a safe final filename. Keep the original extension."""
    suffix = extension if extension.startswith(".") else (f".{extension}" if extension else "")
    text = Path(str(name or "")).stem if name else ""
    text = text.replace("..", " ")
    text = text.replace(":", "-").replace("/", "-").replace("\\", "-")
    text = text.replace("—", "-").replace("–", "-")
    text = re.sub(r'[\\/:*?"<>|]+', " ", text)
    text = re.sub(r"[\x00-\x1f]", " ", text)
    text = re.sub(r"\s+", " ", text).strip(" .")
    text = re.sub(r"[^\w\s\-().'&]+", "", text, flags=re.UNICODE)
    text = re.sub(r"\s+", " ", text).strip(" .")
    if not text or text.casefold() in {"none", "null", "untitled", "new file", "document", "file"}:
        return ""
    words = text.split(" ")
    if len(words) > 10:
        text = " ".join(words[:10])
    if len(text) > 80:
        text = text[:80].rstrip(" .-_")
    return f"{text}{suffix}" if text else ""


def match_existing_category(hint: str, categories: list[dict]) -> str | None:
    """Pick an existing category when the hint clearly refers to one."""
    needle = re.sub(r"\s+", " ", str(hint or "")).casefold().strip()
    if not needle or is_parking_category(needle):
        return None
    exact = []
    partial = []
    for item in categories:
        name = str(item.get("name") or "")
        if not name or is_parking_category(name):
            continue
        key = name.casefold()
        if key == needle:
            return name
        if key in needle or needle in key:
            exact.append(name)
        else:
            tokens = [part for part in re.split(r"\W+", key) if len(part) > 2]
            if tokens and all(token in needle for token in tokens):
                partial.append(name)
    if len(exact) == 1:
        return exact[0]
    if len(partial) == 1:
        return partial[0]
    return None


def fallback_category(kind: str, extension: str = "") -> str:
    if kind == "image":
        return "Photos"
    if kind == "video":
        return "Videos"
    if kind == "audio":
        return "Audio"
    if kind == "sheet":
        return "Spreadsheets"
    if kind == "pdf" or kind == "docx":
        return "Documents"
    if kind == "text":
        code = {
            ".py", ".js", ".ts", ".tsx", ".jsx", ".go", ".rs", ".java", ".c", ".cpp",
            ".h", ".hpp", ".cs", ".php", ".rb", ".swift", ".kt", ".sh", ".zsh", ".sql",
        }
        if extension in code:
            return "Code"
        return "Documents"
    if kind == "archive":
        return "Archives"
    return "Files"


def human_size(num: int) -> str:
    value = float(max(num, 0))
    for unit in ("B", "KB", "MB", "GB", "TB"):
        if value < 1024 or unit == "TB":
            if unit == "B":
                return f"{int(value)} B"
            return f"{value:.1f} {unit}"
        value /= 1024
    return f"{num} B"


def library_label(count: int, nbytes: int) -> str:
    if count == 0:
        return "Empty folder"
    noun = "file" if count == 1 else "files"
    return f"{count} {noun} · {human_size(nbytes)}"


def refuse_library(path: Path) -> str | None:
    try:
        resolved = path.expanduser().resolve()
    except OSError:
        return "That folder is not available."
    if not resolved.exists() or not resolved.is_dir():
        return "That folder is not available."
    app_dir = APP_DIR.resolve()
    if resolved == app_dir or _is_inside(resolved, app_dir):
        return "Choose a library folder outside the Arkiv app."
    home = Path.home().resolve()
    blocked = {
        Path("/"),
        Path("/System"),
        Path("/usr"),
        Path("/bin"),
        Path("/sbin"),
        Path("/Library"),
        Path("/Applications"),
        Path("/private"),
        Path("/Volumes"),
        Path("/Users"),
        home,
        home / "Library",
        home / "AppData",
    }
    # Windows system roots and profile root.
    for letter in "CDEFGHIJKLMNOPQRSTUVWXYZ":
        blocked.add(Path(f"{letter}:/"))
        blocked.add(Path(f"{letter}:/Windows"))
        blocked.add(Path(f"{letter}:/Program Files"))
        blocked.add(Path(f"{letter}:/Program Files (x86)"))
        blocked.add(Path(f"{letter}:/Users"))
    blocked = {item.resolve() for item in blocked if item.exists() or item == home}
    if resolved in blocked:
        return (
            "Choose a specific folder, such as Desktop, Documents, or a project folder. "
            "Arkiv will not take over your whole home folder."
        )
    try:
        home.relative_to(resolved)
        return "That folder is too broad. Choose a folder inside your home directory or on a drive."
    except ValueError:
        return None


def should_skip_name(path: Path) -> bool:
    name = path.name
    if name.startswith("~$") or name.startswith("._") or name.startswith("."):
        return True
    if name.lower() in IGNORED_NAMES:
        return True
    lower = name.lower()
    return lower.endswith(TEMP_SUFFIXES)


def should_skip_dir(path: Path) -> bool:
    name = path.name
    if name.startswith(".") or name.lower() in ALWAYS_SKIP_DIRS or name in ALWAYS_SKIP_DIRS:
        return True
    if name.lower() in {"venv", "env"} and (path / "pyvenv.cfg").exists():
        return True
    return False


def iter_files(root: Path, category_names: set[str], app_dir: Path | None = None) -> list[Path]:
    root = root.resolve()
    app_dir = (app_dir or APP_DIR).resolve()
    categories = {name.casefold() for name in category_names}
    found: list[Path] = []
    for dirpath, dirnames, filenames in os.walk(root, followlinks=False):
        current = Path(dirpath)
        kept: list[str] = []
        for dirname in dirnames:
            child = current / dirname
            if child.is_symlink() or should_skip_dir(child):
                continue
            try:
                resolved = child.resolve()
            except OSError:
                continue
            if resolved == app_dir or _is_inside(resolved, app_dir):
                continue
            if current == root and dirname.casefold() in categories:
                continue
            kept.append(dirname)
        dirnames[:] = kept
        try:
            current_resolved = current.resolve()
        except OSError:
            continue
        if current_resolved == app_dir or _is_inside(current_resolved, app_dir):
            continue
        for filename in filenames:
            path = current / filename
            if path.is_symlink() or not path.is_file() or should_skip_name(path):
                continue
            found.append(path)
    found.sort(key=lambda item: str(item).casefold())
    return found


def _is_inside(path: Path, parent: Path) -> bool:
    try:
        path.relative_to(parent)
    except ValueError:
        return False
    return True


def top_folder(root: Path, path: Path) -> str | None:
    relative = path.resolve().relative_to(root.resolve())
    if len(relative.parts) < 2:
        return None
    return relative.parts[0]


def measure(paths: list[Path]) -> tuple[int, int]:
    total = 0
    for path in paths:
        try:
            total += path.stat().st_size
        except OSError:
            continue
    return len(paths), total


def describe(path: Path, root: Path, *, snippet: bool = False) -> dict:
    extension = path.suffix.lower()
    size = path.stat().st_size
    kind = _kind(extension)
    limit = 160 if snippet else 7000
    text = None
    image = None
    if size <= CONTENT_LIMIT and not snippet:
        text = _read_content(path, kind, limit)
    elif size <= CONTENT_LIMIT and snippet and kind in {"text", "pdf", "docx", "sheet"}:
        text = _read_content(path, kind, limit)
    if kind == "image" and not snippet and size <= IMAGE_LIMIT:
        image = _image_data_url(path)
    relative = str(path.resolve().relative_to(root.resolve()))
    return {
        "name": path.name,
        "relative": relative,
        "extension": extension or "none",
        "size": size,
        "kind": kind,
        "text": (text or "").strip() or None,
        "image": image,
    }


def _kind(extension: str) -> str:
    if extension in IMAGE_EXTENSIONS:
        return "image"
    if extension in PDF_EXTENSIONS:
        return "pdf"
    if extension in DOCX_EXTENSIONS:
        return "docx"
    if extension in SHEET_EXTENSIONS:
        return "sheet"
    if extension in TEXT_EXTENSIONS:
        return "text"
    if extension in {".mp4", ".mov", ".m4v", ".mkv", ".avi", ".webm"}:
        return "video"
    if extension in {".mp3", ".wav", ".m4a", ".aac", ".flac"}:
        return "audio"
    if extension in {".zip", ".rar", ".7z", ".tar", ".gz", ".dmg", ".iso"}:
        return "archive"
    return "binary"


def _read_content(path: Path, kind: str, limit: int) -> str | None:
    try:
        if kind == "text":
            return _read_text(path, limit)
        if kind == "pdf":
            return _read_pdf(path, limit)
        if kind == "docx":
            return _read_docx(path, limit)
        if kind == "sheet":
            return _read_sheet(path, limit)
    except OSError:
        return None
    return None


def _read_text(path: Path, limit: int) -> str:
    with path.open("rb") as handle:
        raw = handle.read(max(limit * 4, 4096))
    return raw.decode("utf-8", errors="replace")[:limit]


def _read_pdf(path: Path, limit: int) -> str | None:
    try:
        from pypdf import PdfReader
    except ImportError:
        return None
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            return None
        parts: list[str] = []
        total = 0
        for page in reader.pages[:8]:
            piece = page.extract_text() or ""
            parts.append(piece)
            total += len(piece)
            if total >= limit:
                break
        return "\n".join(parts)[:limit]
    except Exception:
        return None


def _read_docx(path: Path, limit: int) -> str | None:
    try:
        from docx import Document
    except ImportError:
        return None
    try:
        document = Document(str(path))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        return "\n".join(parts)[:limit]
    except Exception:
        return None


def _read_sheet(path: Path, limit: int) -> str | None:
    try:
        import openpyxl
    except ImportError:
        return None
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return None
    try:
        sheet = workbook.worksheets[0]
        lines: list[str] = []
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if index >= 30:
                break
            cells = [str(cell) for cell in row if cell is not None]
            if cells:
                lines.append(" | ".join(cells[:12]))
        return "\n".join(lines)[:limit]
    except Exception:
        return None
    finally:
        workbook.close()


def _image_data_url(path: Path) -> str | None:
    try:
        from PIL import Image
    except ImportError:
        return None
    try:
        with Image.open(path) as image:
            frame = image.convert("RGB")
            frame.thumbnail((1280, 1280))
            buffer = io.BytesIO()
            frame.save(buffer, format="JPEG", quality=76)
    except Exception:
        return None
    payload = base64.b64encode(buffer.getvalue()).decode("ascii")
    if len(payload) > 4_500_000:
        return None
    return "data:image/jpeg;base64," + payload


def contained(root: Path, path: Path) -> bool:
    try:
        path.resolve().relative_to(root.resolve())
    except (OSError, ValueError):
        return False
    return True


def unique_destination(folder: Path, filename: str) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    candidate = folder / filename
    if not candidate.exists():
        return candidate
    stem = Path(filename).stem
    suffix = Path(filename).suffix
    number = 2
    while True:
        candidate = folder / f"{stem} ({number}){suffix}"
        if not candidate.exists():
            return candidate
        number += 1


def place_file(root: Path, category: str, source: Path, new_name: str | None = None) -> Path:
    folder_name = clean_category(category)
    folder = (root / folder_name).resolve()
    if not _is_inside(folder, root.resolve()) and folder != root.resolve():
        raise ValueError("Category folder would leave the library.")
    filename = new_name or source.name
    filename = clean_filename(filename, source.suffix) or source.name
    destination = unique_destination(folder, filename)
    if source.resolve() == destination.resolve():
        return source
    # Same folder, only a rename.
    if source.resolve().parent == folder and destination.parent == folder:
        source.rename(destination)
        return destination
    shutil.move(str(source), str(destination))
    return destination


def restore_file(root: Path, src: Path, dest: Path) -> Path | None:
    if not dest.exists():
        return None
    parent = src.parent
    try:
        parent_resolved = parent.resolve() if parent.exists() else parent
        root_resolved = root.resolve()
        if parent_resolved != root_resolved and not _is_inside(parent_resolved, root_resolved):
            return None
    except OSError:
        return None
    parent.mkdir(parents=True, exist_ok=True)
    target = src if not src.exists() else unique_destination(parent, src.name)
    shutil.move(str(dest), str(target))
    _remove_if_empty(dest.parent, root)
    return target


def _remove_if_empty(folder: Path, root: Path) -> None:
    try:
        if folder.resolve() == root.resolve():
            return
        folder.rmdir()
    except OSError:
        return
